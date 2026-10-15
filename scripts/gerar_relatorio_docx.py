"""Gera o relatorio .docx a partir dos relatorios em Markdown.

O Markdown e a fonte unica do texto: altere os .md e rode este script de novo.
Suporta o subconjunto de Markdown usado nos relatorios: titulos, paragrafos,
negrito/italico/codigo/links, listas, tabelas, imagens, blocos de codigo e '---'.

Uso (na raiz do repositorio):
    python scripts/gerar_relatorio_docx.py
"""
import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

RELATORIO_PRINCIPAL = 'RELATORIO_TRABALHO_AV1_ML.md'
ANEXOS = ['RELATORIO_PROBLEMA1_COMPETICAO.md']
CAMINHO_SAIDA = os.path.join('outputs', 'Relatorio_Trabalho_Pratico_AV1_ML.docx')

AZUL = RGBColor(0x1F, 0x49, 0x7D)
LARGURA_IMAGEM = Inches(6.3)

# Ordem importa: link e negrito antes de italico (os dois usam '*')
PADRAO_INLINE = re.compile(r'\[(?P<link>[^\]]+)\]\([^)]+\)'
                           r'|\*\*(?P<negrito>.+?)\*\*'
                           r'|`(?P<codigo>[^`]+)`'
                           r'|\*(?P<italico>[^*]+)\*')
PADRAO_IMAGEM = re.compile(r'^!\[(?P<alt>[^\]]*)\]\((?P<caminho>[^)]+)\)$')
PADRAO_TITULO = re.compile(r'^(?P<nivel>#{1,4})\s+(?P<texto>.+)$')
PADRAO_LISTA = re.compile(r'^(?P<marcador>-|\d+\.)\s+(?P<texto>.+)$')


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def adicionar_inline(paragrafo, texto, negrito=False, italico=False):
    """Escreve o texto no paragrafo convertendo **negrito**, *italico*, `codigo` e [links](url)."""
    posicao = 0
    for m in PADRAO_INLINE.finditer(texto):
        if m.start() > posicao:
            run = paragrafo.add_run(texto[posicao:m.start()])
            run.bold, run.italic = negrito, italico
        if m.group('link') is not None:
            adicionar_inline(paragrafo, m.group('link'), negrito, italico)
        elif m.group('negrito') is not None:
            adicionar_inline(paragrafo, m.group('negrito'), True, italico)
        elif m.group('italico') is not None:
            adicionar_inline(paragrafo, m.group('italico'), negrito, True)
        else:
            run = paragrafo.add_run(m.group('codigo'))
            run.bold, run.italic = negrito, italico
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
        posicao = m.end()
    if posicao < len(texto):
        run = paragrafo.add_run(texto[posicao:])
        run.bold, run.italic = negrito, italico


def adicionar_tabela(doc, linhas):
    celulas = [[c.strip() for c in l.strip().strip('|').split('|')] for l in linhas]
    cabecalho, separador, corpo = celulas[0], celulas[1], celulas[2:]
    centralizada = [s.startswith(':') and s.endswith(':') for s in separador]

    tabela = doc.add_table(rows=len(corpo) + 1, cols=len(cabecalho))
    tabela.style = 'Table Grid'
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, linha in enumerate([cabecalho] + corpo):
        for j, texto in enumerate(linha[:len(cabecalho)]):
            cell = tabela.cell(i, j)
            p = cell.paragraphs[0]
            adicionar_inline(p, texto, negrito=(i == 0))
            if centralizada[j]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(9.5)
                if i == 0:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if i == 0:
                set_cell_background(cell, '1F497D')
    doc.add_paragraph()


def adicionar_imagem(doc, alt, caminho):
    if not os.path.exists(caminho):
        print(f"  Aviso: imagem nao encontrada, pulando: {caminho}")
        return
    doc.add_picture(caminho, width=LARGURA_IMAGEM)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if alt:
        legenda = doc.add_paragraph()
        legenda.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = legenda.add_run(alt)
        run.italic = True
        run.font.size = Pt(9)


def adicionar_titulo(doc, texto, nivel):
    titulo = doc.add_heading(level=nivel)
    adicionar_inline(titulo, texto)
    for run in titulo.runs:
        run.font.color.rgb = AZUL


def converter_markdown(doc, caminho_md, capa):
    """Converte um arquivo Markdown. Com capa=True, o primeiro '#' vira o titulo da capa
    e o primeiro '---' vira a quebra de pagina depois dela."""
    with open(caminho_md, encoding='utf-8') as f:
        linhas = f.read().splitlines()

    titulo_feito = False
    paragrafo = []

    def fechar_paragrafo():
        if paragrafo:
            p = doc.add_paragraph()
            adicionar_inline(p, ' '.join(paragrafo))
            if capa and not quebra_feita[0]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragrafo.clear()

    quebra_feita = [not capa]
    i = 0
    while i < len(linhas):
        linha = linhas[i].rstrip()
        texto = linha.strip()

        if not texto:
            fechar_paragrafo()
        elif texto.startswith('```'):
            fechar_paragrafo()
            i += 1
            codigo = []
            while i < len(linhas) and not linhas[i].strip().startswith('```'):
                codigo.append(linhas[i])
                i += 1
            p = doc.add_paragraph()
            run = p.add_run('\n'.join(codigo))
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            p.paragraph_format.left_indent = Inches(0.3)
        elif texto == '---':
            fechar_paragrafo()
            if not quebra_feita[0]:
                doc.add_page_break()
                quebra_feita[0] = True
        elif PADRAO_TITULO.match(texto):
            fechar_paragrafo()
            m = PADRAO_TITULO.match(texto)
            nivel = len(m.group('nivel'))
            if capa and not titulo_feito and nivel == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(m.group('texto'))
                run.bold = True
                run.font.size = Pt(22)
                run.font.color.rgb = AZUL
                titulo_feito = True
            else:
                adicionar_titulo(doc, m.group('texto'), nivel)
        elif texto.startswith('|'):
            fechar_paragrafo()
            tabela = []
            while i < len(linhas) and linhas[i].strip().startswith('|'):
                tabela.append(linhas[i])
                i += 1
            adicionar_tabela(doc, tabela)
            continue
        elif PADRAO_IMAGEM.match(texto):
            fechar_paragrafo()
            m = PADRAO_IMAGEM.match(texto)
            adicionar_imagem(doc, m.group('alt'), m.group('caminho'))
        elif PADRAO_LISTA.match(texto):
            fechar_paragrafo()
            m = PADRAO_LISTA.match(texto)
            recuo = (len(linha) - len(linha.lstrip())) // 2
            if m.group('marcador') == '-':
                p = doc.add_paragraph(style='List Bullet 2' if recuo else 'List Bullet')
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.25 + 0.25 * recuo)
                p.add_run(m.group('marcador') + ' ')
            adicionar_inline(p, m.group('texto'))
        else:
            paragrafo.append(texto)
        i += 1
    fechar_paragrafo()


def gerar_relatorio_docx():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    font = doc.styles['Normal'].font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    converter_markdown(doc, RELATORIO_PRINCIPAL, capa=True)
    for anexo in ANEXOS:
        doc.add_page_break()
        converter_markdown(doc, anexo, capa=False)

    os.makedirs(os.path.dirname(CAMINHO_SAIDA), exist_ok=True)
    doc.save(CAMINHO_SAIDA)
    print(f"Relatorio salvo em: {CAMINHO_SAIDA}")


if __name__ == '__main__':
    gerar_relatorio_docx()
